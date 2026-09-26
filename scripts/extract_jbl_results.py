#!/usr/bin/env python3
"""Read-only extraction of Kurt's results workbooks. Never sum partial rosters."""
import argparse
import datetime
import hashlib
import json
import re
from pathlib import Path

import openpyxl


def team_id(value):
    text = str(value or '').upper()
    aliases = {'CELTICS': 'central-celtics', 'KNIGHTS': 'urban-knights',
               'EAGLES': 'upper-room-eagles', 'REBELS': 'uwi-running-rebels',
               'FLAMES': 'portmore-flames', 'SLAYERS': 'st-georges-slayers',
               'SPARTANS': 'spanish-town-spartans', 'RAPTORS': 'rae-town-raptors',
               'WIZZARDS': 'tivoli-wizards', 'WIZARDS': 'tivoli-wizards'}
    for word, identity in aliases.items():
        if word in text:
            return identity
    # One report says Tivoli Warriors: ambiguous, not silently mapped.
    if 'WARRIORS' in text and 'TIVOLI' not in text:
        return 'mbbgc-warriors'
    return None


def number(value):
    if isinstance(value, (int, float)) and value >= 0 and int(value) == value:
        return int(value)
    return None


def date_value(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    text = re.sub(r'^(DATE\s*:?)', '', str(value or '').strip(), flags=re.I)
    text = re.sub(r'(\d)(st|nd|rd|th)', r'\1', text, flags=re.I)
    text = ' '.join(text.replace(',', ' ').split()).upper()
    for fmt in ('%B %d %Y', '%b %d %Y', '%d %B %Y', '%d %b %Y'):
        try:
            return datetime.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def extract(folder):
    games, issues, files = {}, [], []
    for path in sorted(Path(folder).glob('*.xlsx')):
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        files.append({'name': path.name, 'sha256': sha})
        book = openpyxl.load_workbook(path, data_only=True)
        for sheet in book:
            if sheet.max_row < 4:
                continue
            date = date_value(sheet['B1'].value)
            if not date:
                issues.append({'file': path.name, 'sheet': sheet.title, 'reason': 'Undated worksheet excluded'})
                continue
            for a, b in ((2, 11), (23, 33)):
                sides = []
                for row in (a, b):
                    raw = sheet.cell(row, 1).value
                    if not team_id(raw):
                        sides = []
                        issues.append({'file': path.name, 'sheet': sheet.title,
                                       'cell': f'A{row}', 'reason': 'Unresolved team', 'value': raw})
                        break
                    final = number(sheet.cell(row + 2, 8).value)
                    quarters = [number(sheet.cell(row + 2, c).value) for c in range(3, 7)]
                    overtime = number(sheet.cell(row + 2, 7).value)
                    headers = {str(sheet.cell(row + 3, c).value).strip().upper(): c for c in range(3, 8)}
                    players = []
                    # Read only named lines, not percentages or inherited empty formula rows.
                    end = b if row == a else (22 if b == 11 else sheet.max_row + 1)
                    for r in range(row + 4, end):
                        name = sheet.cell(r, 2).value
                        if not isinstance(name, str) or not name.strip():
                            continue
                        if name.strip().upper() in ('NAME', 'QUARTERS', 'QUARTER SCORES'):
                            continue
                        stats = {key.lower(): number(sheet.cell(r, col).value) for key, col in headers.items()
                                 if key in ('POINTS', 'REBOUNDS', 'ASSISTS', 'STEALS', 'BLOCKS')}
                        if stats.get('points') is None:
                            continue
                        players.append({'name': name.strip(), **stats})
                    sides.append({'teamId': team_id(raw), 'score': final, 'quarters': quarters,
                                  'overtime': overtime, 'players': players})
                if len(sides) != 2 or any(s['score'] is None for s in sides):
                    continue
                if sides[0]['score'] == sides[1]['score']:
                    issues.append({'file': path.name, 'sheet': sheet.title, 'row': a, 'reason': 'Tied or unplayed final excluded'})
                    continue
                key = date + '-' + '-vs-'.join(sorted(s['teamId'] for s in sides))
                periods_valid = all(all(q is not None for q in s['quarters']) and
                                    sum(s['quarters']) + (s['overtime'] or 0) == s['score'] for s in sides)
                if not periods_valid:
                    issues.append({'file': path.name, 'sheet': sheet.title, 'row': a, 'reason': 'Quarter total does not match final; omit breakdown'})
                source = {'file': path.name, 'sha256': sha, 'sheet': sheet.title, 'range': f'A{a}:H{min(42,b+8)}'}
                game = {'gameId': 'jbl-' + key, 'date': date, 'sides': sides,
                        'periodsVerified': periods_valid, 'sources': [source]}
                if key in games:
                    prior = games[key]
                    if prior['sides'] != sides:
                        issues.append({'file': path.name, 'sheet': sheet.title, 'reason': 'Conflicting duplicate', 'gameId': game['gameId']})
                        prior['conflict'] = True
                    prior['sources'].append(source)
                else:
                    games[key] = game
    return {'files': files, 'games': sorted([g for g in games.values() if not g.get('conflict')], key=lambda g: g['date']),
            'issues': issues, 'coverage': 'Selected game reports, not a complete season ledger. Player lines are selected performers, not full rosters.'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('folder')
    args = parser.parse_args()
    print(json.dumps(extract(args.folder), ensure_ascii=False, indent=2))
